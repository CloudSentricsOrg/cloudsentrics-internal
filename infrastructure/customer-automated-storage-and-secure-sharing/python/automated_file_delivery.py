import os
import re
import csv
import io
import json
import hashlib
import base64
import time
import urllib.parse
from datetime import datetime, timedelta, timezone

import boto3
from botocore.exceptions import ClientError

# ========= Environment =========
MAILHUB_ACCOUNT_ID = os.getenv("MAILHUB_ACCOUNT_ID")
MAILHUB_ROLE_NAME  = os.getenv("MAILHUB_ROLE_NAME", "centralized-ses-role")
MAILHUB_REGION     = os.getenv("MAILHUB_REGION", "us-east-1")
FROM_EMAIL         = os.environ["FROM_EMAIL"]
STORAGE_BUCKET     = os.environ["STORAGE_BUCKET"]
OTP_TABLE          = os.environ["OTP_TABLE"]
PUBLIC_BASE_URL    = os.environ["PUBLIC_BASE_URL"]
DELIVERY_TABLE     = os.environ["DELIVERY_TABLE"]
ROOT_FOLDER        = os.getenv("ROOT_FOLDER", "records")
INVITE_TTL_HOURS   = int(os.getenv("INVITE_TTL_HOURS", "48"))

# ========= Clients =========
s3     = boto3.client("s3")
sts    = boto3.client("sts")
dynamo = boto3.client("dynamodb")

# ========= Helpers =========
def _now():
    return datetime.now(timezone.utc)

def _sha256(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()

def _rand_token_b64url(nbytes: int = 32) -> str:
    return base64.urlsafe_b64encode(os.urandom(nbytes)).rstrip(b"=").decode("ascii")

def _assume_mailhub():
    role_arn = f"arn:aws:iam::{MAILHUB_ACCOUNT_ID}:role/{MAILHUB_ROLE_NAME}"
    resp = sts.assume_role(RoleArn=role_arn, RoleSessionName="cs-file-delivery-ses")
    return resp["Credentials"]

def _mailhub_ses():
    c = _assume_mailhub()
    return boto3.client(
        "ses",
        region_name=MAILHUB_REGION,
        aws_access_key_id=c["AccessKeyId"],
        aws_secret_access_key=c["SecretAccessKey"],
        aws_session_token=c["SessionToken"],
    )

def _send_with_retry(fn, attempts=3, base_delay=0.4):
    for i in range(1, attempts + 1):
        try:
            return fn()
        except ClientError as e:
            print(f"[retry] send attempt {i} failed:", getattr(e, "response", {}))
            if i == attempts:
                raise
            time.sleep(base_delay * (2 ** (i - 1)))

def _get_bucket_org(bucket: str) -> dict:
    org = {"name": "Your Organization", "email": "", "phone": "", "address": ""}
    try:
        btags = s3.get_bucket_tagging(Bucket=bucket).get("TagSet", [])
        for t in btags:
            k = t["Key"].lower()
            if k == "organization":
                org["name"] = t["Value"] or org["name"]
            elif k == "organizationemail":
                org["email"] = t["Value"]
            elif k == "organizationphone":
                org["phone"] = t["Value"]
            elif k == "organizationaddress":
                org["address"] = t["Value"]
    except Exception as e:
        print(f"No/err bucket tags for {bucket}: {e}")
    return org

# ========= Flexible Column Matching =========
_EMAIL_FIELDS = {"recipientemail", "email", "email address", "emailaddress", "email_address",
                 "patient email", "patientemail", "patient_email", "parent email", "parentemail",
                 "parent_email", "customer_email", "client_email"}
_PHONE_FIELDS = {"recipientphone", "phone", "phone number", "phonenumber", "phone_number",
                 "patient phone", "patientphone", "patient_phone", "parent phone", "parentphone",
                 "parent_phone", "mobile", "cell", "whatsapp"}
_NAME_FIELDS = {"recipientname", "name", "patient name", "patientname", "parent name", "parentname",
                "parent_name", "full name", "fullname", "pat_name", "patient_name", "client_name",
                "customer_name"}
_FIRST_NAME_FIELDS = {"first name", "firstname", "fname", "pat_first_name", "first"}
_LAST_NAME_FIELDS = {"last name", "lastname", "lname", "pat_last_name", "last"}
_ID_FIELDS = {"studentid", "student_id", "student id", "patientid", "patient_id", "patient id",
              "id", "identifier", "mrn", "case_id", "caseid", "case id", "client_id", "clientid",
              "policy_id", "policyid", "account_id", "accountid"}
_RESEND_FIELDS = {"resend"}

_CONTROL_FIELDS = _EMAIL_FIELDS | _PHONE_FIELDS | _NAME_FIELDS | _FIRST_NAME_FIELDS | _LAST_NAME_FIELDS | _ID_FIELDS | _RESEND_FIELDS

def _extract_entry(row: dict) -> dict:
    lower_row = {k.lower().strip(): v for k, v in row.items()}

    email = ""
    for f in _EMAIL_FIELDS:
        if f in lower_row and lower_row[f]:
            email = lower_row[f].strip()
            break

    phone = ""
    for f in _PHONE_FIELDS:
        if f in lower_row and lower_row[f]:
            phone = lower_row[f].strip()
            break

    name = ""
    for f in _NAME_FIELDS:
        if f in lower_row and lower_row[f]:
            name = lower_row[f].strip()
            break
    if not name:
        first = last = ""
        for f in _FIRST_NAME_FIELDS:
            if f in lower_row and lower_row[f]:
                first = lower_row[f].strip()
                break
        for f in _LAST_NAME_FIELDS:
            if f in lower_row and lower_row[f]:
                last = lower_row[f].strip()
                break
        name = f"{first} {last}".strip()

    identifier = ""
    for f in _ID_FIELDS:
        if f in lower_row and lower_row[f]:
            identifier = lower_row[f].strip()
            break

    resend = ""
    for f in _RESEND_FIELDS:
        if f in lower_row and lower_row[f]:
            resend = lower_row[f].strip()
            break

    # Sensitive data — everything not a control field
    sensitive_data = {}
    for orig_key, value in row.items():
        if orig_key.lower().strip() not in _CONTROL_FIELDS and value and str(value).strip():
            sensitive_data[orig_key] = str(value).strip()

    return {"email": email, "phone": phone, "name": name, "identifier": identifier,
            "resend": resend, "sensitive_data": sensitive_data}

# ========= File Reading =========
def _read_spreadsheet(bucket: str, key: str) -> list[dict]:
    try:
        obj = s3.get_object(Bucket=bucket, Key=key)
        raw = obj["Body"].read()

        # Excel (.xlsx)
        if key.lower().endswith(".xlsx"):
            return _parse_excel(raw)

        body = raw.decode("utf-8-sig")

        # JSON
        try:
            data = json.loads(body)
            if isinstance(data, list):
                return [_extract_entry(r) for r in data if isinstance(r, dict)]
        except (json.JSONDecodeError, ValueError):
            pass

        # CSV
        reader = csv.DictReader(io.StringIO(body))
        return [_extract_entry(row) for row in reader]
    except Exception as e:
        print(f"Error reading s3://{bucket}/{key}: {e}")
        return []

def _parse_excel(raw_bytes: bytes) -> list[dict]:
    """Parse Excel .xlsx file into list of dicts."""
    from openpyxl import load_workbook
    wb = load_workbook(filename=io.BytesIO(raw_bytes), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h).strip() if h else "" for h in rows[0]]
    entries = []
    for row in rows[1:]:
        row_dict = {headers[i]: str(v).strip() if v is not None else "" for i, v in enumerate(row) if i < len(headers)}
        entries.append(_extract_entry(row_dict))
    wb.close()
    return entries

# ========= Storage Bucket File Matching =========
def _find_files_for_id(identifier: str) -> list[str]:
    """Find all files in the storage bucket matching the identifier."""
    prefix = f"{ROOT_FOLDER}/{identifier}"
    try:
        resp = s3.list_objects_v2(Bucket=STORAGE_BUCKET, Prefix=prefix)
        files = []
        for obj in resp.get("Contents", []):
            key = obj["Key"]
            if not key.endswith("/"):
                files.append(key)
        return files
    except Exception as e:
        print(f"Error listing files for {identifier}: {e}")
        return []

# ========= Delivery Tracking =========
def _is_already_delivered(identifier: str, file_key: str) -> bool:
    tracking_key = f"{identifier}|{file_key}"
    try:
        resp = dynamo.get_item(
            TableName=DELIVERY_TABLE,
            Key={"trackingKey": {"S": tracking_key}}
        )
        return "Item" in resp
    except Exception as e:
        print(f"Delivery check failed: {e}")
        return False

def _mark_delivered(identifier: str, file_key: str, recipient_email: str):
    tracking_key = f"{identifier}|{file_key}"
    try:
        dynamo.put_item(
            TableName=DELIVERY_TABLE,
            Item={
                "trackingKey": {"S": tracking_key},
                "identifier": {"S": identifier},
                "fileKey": {"S": file_key},
                "recipientEmail": {"S": recipient_email},
                "deliveredAt": {"S": _now().isoformat()},
            }
        )
    except Exception as e:
        print(f"Mark delivered failed: {e}")

# ========= OTP Invite Creation =========
def _ddb_put_invite(request_id: str, item: dict):
    attrs = {
        "requestId":       {"S": request_id},
        "bucket":          {"S": item["bucket"]},
        "key":             {"S": item["key"]},
        "recipientEmail":  {"S": item["recipientEmail"]},
        "recipientPhone":  {"S": item["recipientPhone"]},
        "orgName":         {"S": item.get("orgName", "Your Organization")},
        "status":          {"S": "INVITED"},
        "inviteExpiresAt": {"N": str(item["inviteExpiresAt"])},
        "linkTokenHash":   {"S": item["linkTokenHash"]},
        "sendWelcome":     {"BOOL": False},
        "shareFile":       {"BOOL": True},
    }
    if item.get("recipientName"):
        attrs["recipientName"] = {"S": item["recipientName"]}
    if item.get("sensitiveData"):
        attrs["sensitiveData"] = {"S": json.dumps(item["sensitiveData"])}

    dynamo.put_item(
        TableName=OTP_TABLE,
        Item=attrs,
        ConditionExpression="attribute_not_exists(requestId)"
    )

# ========= Welcome Email =========
def _send_delivery_email(ses, to_email: str, verify_url: str, org_name: str, recipient_name: str | None, file_count: int):
    display = (org_name or "Cloud Sentrics").strip()
    source_header = f"{display} <{FROM_EMAIL}>"
    greeting = f"Hello {recipient_name}," if recipient_name else "Hello,"
    file_word = "document" if file_count == 1 else "documents"

    html = f"""
    <html><body style="font-family:'Segoe UI',Roboto,Arial,sans-serif;line-height:1.7;color:#1e293b;margin:0;padding:0;background:#f1f5f9;">
      <div style="max-width:600px;margin:40px auto;background:#ffffff;border-radius:16px;overflow:hidden;box-shadow:0 8px 30px rgba(0,0,0,.1);">
        <div style="background:linear-gradient(135deg,#0f766e,#14b8a6,#2dd4bf);padding:36px 32px;text-align:center;">
          <div style="font-size:36px;margin-bottom:12px;">📤</div>
          <h1 style="margin:0;color:#ffffff;font-size:26px;font-weight:700;">A secure file is available</h1>
        </div>
        <div style="padding:36px 32px 28px;">
          <p style="font-size:18px;margin-top:0;color:#0f766e;font-weight:600;">{greeting}</p>
          <p style="font-size:17px;color:#0f766e;font-weight:600;">A file has been securely shared with you by <strong style="color:#1e293b;">{org_name}</strong>.</p>
          <p style="font-size:16px;color:#0f766e;font-weight:500;">Click the button below to start verification. We’ll send a one-time code to your WhatsApp.</p>
          <div style="text-align:center;margin:28px 0;">
            <a href="{verify_url}" target="_blank"
               style="background:linear-gradient(135deg,#0f766e,#14b8a6);color:#ffffff;padding:14px 40px;border-radius:50px;text-decoration:none;font-weight:700;font-size:15px;display:inline-block;box-shadow:0 4px 14px rgba(20,184,166,.4);letter-spacing:0.3px;">
              Verify &amp; Continue →
            </a>
          </div>
        </div>
        <div style="background:#f8fafc;padding:20px 32px;text-align:center;border-top:1px solid #e2e8f0;">
          <p style="margin:0;font-size:13px;font-weight:600;color:#14b8a6;">
            🛡️ Secure delivery powered by Cloud Sentrics
          </p>
        </div>
      </div>
    </body></html>
    """

    _send_with_retry(lambda: ses.send_email(
        Source=source_header,
        Destination={"ToAddresses": [to_email]},
        Message={
            "Subject": {"Data": f"A file has been securely shared by {org_name}"},
            "Body": {"Html": {"Data": html}}
        },
    ))


# ========= Main handler =========
def lambda_handler(event, context):
    # Get the uploaded spreadsheet info from S3 event
    record = event["Records"][0]
    delivery_bucket = record["s3"]["bucket"]["name"]
    spreadsheet_key = urllib.parse.unquote_plus(record["s3"]["object"]["key"])

    print(f"Processing spreadsheet: {spreadsheet_key}")

    ses = _mailhub_ses()
    org = _get_bucket_org(STORAGE_BUCKET)
    org_name = org["name"]

    entries = _read_spreadsheet(delivery_bucket, spreadsheet_key)
    if not entries:
        print("No entries found in spreadsheet")
        return {"statusCode": 200, "delivered": 0, "skipped": 0}

    EMAIL_REGEX = r"[^@]+@[^@]+\.[^@]+"
    PHONE_REGEX = r"^\+\d{6,15}$"

    delivered = 0
    skipped = 0

    for entry in entries:
        email = entry["email"]
        phone = entry["phone"]
        name = entry["name"]
        identifier = entry["identifier"]
        resend_raw = entry.get("resend", "").strip()
        resend = True if resend_raw.lower() == "true" else (resend_raw if resend_raw else False)
        sensitive_data = entry["sensitive_data"]

        if not email or not re.match(EMAIL_REGEX, email):
            print(f"Skipping row: invalid or missing email")
            continue
        if not phone or not re.match(PHONE_REGEX, phone):
            print(f"Skipping row: invalid or missing phone")
            continue
        if not identifier:
            print(f"Skipping row: no identifier provided")
            continue

        # Find files in storage bucket
        files = _find_files_for_id(identifier)
        if not files:
            print(f"Skipping {identifier}: no files found in storage bucket")
            skipped += 1
            continue

        # Filter based on Resend column
        if resend:
            if resend is True:
                pass  # resend=true: send all files
            else:
                # resend=specific file name
                files = [f for f in files if f.endswith(resend)]
        else:
            # Normal delivery — only new files
            files = [f for f in files if not _is_already_delivered(identifier, f)]

        if not files:
            print(f"Skipping {identifier}: already delivered, no resend specified")
            skipped += 1
            continue

        # Create OTP invite for each file
        for file_key in files:
            request_id = hashlib.sha1(f"{file_key}|{email}|{_now().timestamp()}".encode()).hexdigest()[:20]
            invite_exp = int((_now() + timedelta(hours=INVITE_TTL_HOURS)).timestamp())
            link_token = _rand_token_b64url(32)

            try:
                _ddb_put_invite(request_id, {
                    "bucket": STORAGE_BUCKET,
                    "key": file_key,
                    "recipientEmail": email,
                    "recipientPhone": phone,
                    "recipientName": name,
                    "orgName": org_name,
                    "inviteExpiresAt": invite_exp,
                    "linkTokenHash": _sha256(link_token),
                    "sensitiveData": sensitive_data if sensitive_data else None,
                })
            except Exception as e:
                print(f"DDB put invite failed for {file_key}: {e}")
                continue

            verify_url = f"{PUBLIC_BASE_URL}/start?rid={urllib.parse.quote(request_id)}&tok={urllib.parse.quote(link_token)}"
            try:
                _send_delivery_email(ses, email, verify_url, org_name, name, len(files))
                if not resend:
                    _mark_delivered(identifier, file_key, email)
                delivered += 1
                print(f"Delivered: {file_key} to {email} (rid={request_id})")
            except Exception as e:
                print(f"Email failed for {file_key}: {e}")

    print(f"File delivery complete: {delivered} delivered, {skipped} skipped")
    return {"statusCode": 200, "delivered": delivered, "skipped": skipped}
